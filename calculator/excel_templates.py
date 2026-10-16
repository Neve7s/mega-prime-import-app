"""Plantilla mínima para importar gastos de servicios y productos."""

from io import BytesIO
from zipfile import BadZipFile

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils.exceptions import InvalidFileException


GASTOS_PLANTILLA = {
    "Flete internacional": "flete_intl",
    "Flete local": "flete_local",
    "ISD / comisión bancaria": "comision_banco",
    "Bodegaje": "bodegaje",
    "Devolución": "devolucion",
    "Administración": "gasto_admin",
    "Descarga": "descarga",
    "Agente": "agente",
    "Otros por unidad": "otros_unit",
}

GASTOS_VACIOS = {key: "0" for key in GASTOS_PLANTILLA.values()}


def build_plantilla_xlsx():
    workbook = Workbook()
    gastos = workbook.active
    gastos.title = "Gastos"
    gastos.append(["Servicio", "Valor USD"])
    for label in GASTOS_PLANTILLA:
        gastos.append([label, None])
    gastos.freeze_panes = "A2"
    gastos.column_dimensions["A"].width = 32
    gastos.column_dimensions["B"].width = 18

    productos = workbook.create_sheet("Productos")
    productos.append(["Producto", "Cantidad", "Precio FOB unitario USD"])
    productos.append([None, None, None])
    productos.freeze_panes = "A2"
    productos.column_dimensions["A"].width = 38
    productos.column_dimensions["B"].width = 16
    productos.column_dimensions["C"].width = 26

    for sheet in (gastos, productos):
        for cell in sheet[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="0A0A0B")
        sheet.auto_filter.ref = sheet.dimensions

    output = BytesIO()
    workbook.save(output)
    return output.getvalue()


def read_plantilla_xlsx(uploaded_file):
    try:
        workbook = load_workbook(uploaded_file, read_only=True, data_only=True)
    except (BadZipFile, InvalidFileException, OSError, ValueError) as exc:
        raise ValueError("No se pudo leer el archivo Excel. Use la plantilla descargada.") from exc

    if not {"Gastos", "Productos"}.issubset(workbook.sheetnames):
        raise ValueError("El archivo debe incluir las hojas Gastos y Productos.")

    def number(value, label):
        if value is None or value == "":
            return "0"
        try:
            result = float(value)
        except (TypeError, ValueError):
            raise ValueError(f"El valor de {label} debe ser numérico.") from None
        if not (result >= 0 and result < float("inf")):
            raise ValueError(f"El valor de {label} debe ser un número positivo.")
        return str(value).strip()

    expenses = dict(GASTOS_VACIOS)
    sheet = workbook["Gastos"]
    if [sheet.cell(1, col).value for col in (1, 2)] != ["Servicio", "Valor USD"]:
        raise ValueError("La hoja Gastos debe conservar los encabezados de la plantilla.")
    if sheet.max_row > 20:
        raise ValueError("La hoja Gastos contiene demasiadas filas.")
    label_to_key = {label.casefold(): key for label, key in GASTOS_PLANTILLA.items()}
    for row in sheet.iter_rows(min_row=2, max_row=sheet.max_row, values_only=True):
        label = str(row[0] or "").strip()
        if not label:
            continue
        key = label_to_key.get(label.casefold())
        if key is None:
            raise ValueError(f"Servicio no reconocido en la plantilla: {label}.")
        expenses[key] = number(row[1], label)

    sheet = workbook["Productos"]
    if [sheet.cell(1, col).value for col in (1, 2, 3)] != [
        "Producto", "Cantidad", "Precio FOB unitario USD"
    ]:
        raise ValueError("La hoja Productos debe conservar los encabezados de la plantilla.")
    if sheet.max_row > 502:
        raise ValueError("La plantilla admite hasta 500 productos.")
    items = []
    for row_number, row in enumerate(
            sheet.iter_rows(min_row=2, max_row=sheet.max_row, values_only=True), start=2):
        name, quantity, price = (row + (None, None, None))[:3]
        if all(value is None or str(value).strip() == "" for value in (name, quantity, price)):
            continue
        if not str(name or "").strip():
            raise ValueError(f"Falta el nombre del producto en la fila {row_number}.")
        if quantity is None or quantity == "" or price is None or price == "":
            raise ValueError(f"Complete cantidad y precio del producto en la fila {row_number}.")
        items.append({
            "nombre": str(name).strip()[:200],
            "cant": number(quantity, f"cantidad, fila {row_number}"),
            "valor_unit": number(price, f"precio, fila {row_number}"),
            "compet1": "35",
            "compet2": "35",
            "compet3": "35",
            "seguro_pct": "",
            "advalorem_pct": "",
            "fondinfa_pct": "",
            "iva_pct": "",
        })
        if len(items) > 500:
            raise ValueError("La plantilla admite hasta 500 productos.")

    workbook.close()
    return {"gastos": expenses, "items": items}
