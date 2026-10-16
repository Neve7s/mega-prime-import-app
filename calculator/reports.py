"""Reporte de embarque en .xlsx presentable (openpyxl con estilos corporativos).

Orden: 1) Detalle por producto, 2) Resumen general.
"""

from datetime import datetime
from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

NEGRO = "0A0A0B"
DORADO = "D4AF37"
GRIS_OSCURO = "3F3F46"
GRIS_CLARO = "F4F4F5"
BANDA = "F7F7F9"
BLANCO = "FFFFFF"
AMBAR = "FFF7E0"
VERDE = "047857"

MONEY_FMT = '"USD" #,##0.00'
PCT_FMT = '0.00%'

FILL_NEGRO = PatternFill("solid", fgColor=NEGRO)
FILL_GRIS = PatternFill("solid", fgColor=GRIS_CLARO)
FILL_BANDA = PatternFill("solid", fgColor=BANDA)
FILL_AMBAR = PatternFill("solid", fgColor=AMBAR)
FONT_TITULO = Font(name="Calibri", size=16, bold=True, color=DORADO)
FONT_SUB = Font(name="Calibri", size=10, italic=True, color=GRIS_OSCURO)
FONT_SEC = Font(name="Calibri", size=11, bold=True, color=DORADO)
FONT_SUBSEC = Font(name="Calibri", size=9, bold=True, color=NEGRO)
FONT_HEAD = Font(name="Calibri", size=9, bold=True, color=BLANCO)
FONT_BODY = Font(name="Calibri", size=9, color=NEGRO)
FONT_BOLD = Font(name="Calibri", size=9, bold=True, color=NEGRO)
FONT_VERDE = Font(name="Calibri", size=9, bold=True, color=VERDE)
FONT_DORADO = Font(name="Calibri", size=10, bold=True, color=DORADO)
THIN = Side(style="thin", color="D9D9DE")
MEDIUM = Side(style="medium", color=NEGRO)
BORDER_ALL = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)
LEFT = Alignment(horizontal="left", vertical="center", wrap_text=True)
RIGHT = Alignment(horizontal="right", vertical="center")

DETALLE_COLS = [
    ("#", 4, None),
    ("Producto", 34, "left"),
    ("Cant", 10, "num"),
    ("V. Unit FOB", 12, "money"),
    ("FOB", 13, "money"),
    ("Ratio %", 9, "pct"),
    ("AdV %", 8, "pct"),
    ("Seg %", 8, "pct"),
    ("Fon %", 8, "pct"),
    ("IVA %", 8, "pct"),
    ("Ad-Valorem", 12, "money"),
    ("Flete Intl", 12, "money"),
    ("Seguro", 11, "money"),
    ("Fondinfa", 11, "money"),
    ("Adicionales", 12, "money"),
    ("V. Total", 13, "money"),
    ("Costo Unit", 12, "money"),
    ("May. recargo %", 12, "pct"),
    ("PVP May. s/IGV", 14, "money"),
    ("PVP May. c/IGV", 14, "money"),
    ("Frec. recargo %", 12, "pct"),
    ("PVP Frec. s/IGV", 14, "money"),
    ("PVP Frec. c/IGV", 14, "money"),
    ("Event. recargo %", 12, "pct"),
    ("PVP Event. s/IGV", 14, "money"),
    ("PVP Event. c/IGV", 14, "money"),
    ("Régimen", 10, "center"),
]


def _section(ws, row, text, ncols):
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=ncols)
    c = ws.cell(row=row, column=1, value=text)
    c.font = FONT_SEC
    c.fill = FILL_NEGRO
    c.alignment = Alignment(horizontal="left", vertical="center")
    ws.row_dimensions[row].height = 22
    for col in range(2, ncols + 1):
        ws.cell(row=row, column=col).fill = FILL_NEGRO
    return row + 1


def _marco(ws, min_row, max_row, min_col, max_col):
    """Marco exterior grueso para encuadrar una tabla."""
    for col in range(min_col, max_col + 1):
        ws.cell(row=min_row, column=col).border = Border(
            left=ws.cell(row=min_row, column=col).border.left,
            right=ws.cell(row=min_row, column=col).border.right,
            top=MEDIUM, bottom=ws.cell(row=min_row, column=col).border.bottom)
        ws.cell(row=max_row, column=col).border = Border(
            left=ws.cell(row=max_row, column=col).border.left,
            right=ws.cell(row=max_row, column=col).border.right,
            top=ws.cell(row=max_row, column=col).border.top, bottom=MEDIUM)
    for row in range(min_row, max_row + 1):
        ws.cell(row=row, column=min_col).border = Border(
            left=MEDIUM, right=ws.cell(row=row, column=min_col).border.right,
            top=ws.cell(row=row, column=min_col).border.top,
            bottom=ws.cell(row=row, column=min_col).border.bottom)
        ws.cell(row=row, column=max_col).border = Border(
            left=ws.cell(row=row, column=max_col).border.left, right=MEDIUM,
            top=ws.cell(row=row, column=max_col).border.top,
            bottom=ws.cell(row=row, column=max_col).border.bottom)


def _kv_row(ws, row, label, value, bold=False, fill=None, ncols=3):
    lc = ws.cell(row=row, column=1, value=label)
    lc.font = FONT_BOLD if bold else FONT_BODY
    lc.border = BORDER_ALL
    lc.alignment = LEFT
    if fill:
        lc.fill = fill
    vc = ws.cell(row=row, column=2, value=value)
    vc.number_format = MONEY_FMT
    vc.font = FONT_BOLD if bold else FONT_BODY
    vc.alignment = RIGHT
    vc.border = BORDER_ALL
    if fill:
        vc.fill = fill
    return row + 1


def _subsec(ws, row, text, ncols=3):
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=ncols)
    c = ws.cell(row=row, column=1, value=text)
    c.font = FONT_SUBSEC
    c.fill = FILL_GRIS
    c.alignment = Alignment(horizontal="left", vertical="center")
    for col in range(2, ncols + 1):
        cell = ws.cell(row=row, column=col)
        cell.fill = FILL_GRIS
        cell.border = BORDER_ALL
    c.border = BORDER_ALL
    ws.row_dimensions[row].height = 18
    return row + 1


def build_reporte_xlsx(calc, tasas, gastos, embarque_nombre="Embarque"):
    """calc = resultado de calculate_shipment. Devuelve bytes del .xlsx."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Reporte"
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.oddHeader.center.text = "MEGA PRIME IMPORT — Reporte de embarque"
    ws.oddHeader.center.font = "Calibri,Regular"
    ws.oddHeader.center.size = 9

    R = calc["resumen"]
    T = {k: float(v) for k, v in tasas.items()}
    fecha = datetime.now().strftime("%d/%m/%Y %H:%M")

    # ══ Portada ══
    nd = len(DETALLE_COLS)
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=nd)
    t = ws.cell(row=1, column=1, value="MEGA PRIME IMPORT")
    t.font = FONT_TITULO
    t.fill = FILL_NEGRO
    t.alignment = Alignment(horizontal="left", vertical="center")
    t.border = Border(bottom=Side(style="medium", color=DORADO))
    ws.row_dimensions[1].height = 28
    for col in range(2, nd + 1):
        cell = ws.cell(row=1, column=col)
        cell.fill = FILL_NEGRO
        cell.border = Border(bottom=Side(style="medium", color=DORADO))

    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=nd)
    s = ws.cell(row=2, column=1,
                value="Reporte de embarque — %s · %s" % (embarque_nombre, fecha))
    s.font = FONT_SUB
    ws.row_dimensions[2].height = 16

    ws.merge_cells(start_row=3, start_column=1, end_row=3, end_column=nd)
    c3 = ws.cell(row=3, column=1, value="Cantidad total: %s · %d productos" % (
        R["total_cant"], len(calc["items"])))
    c3.font = FONT_SUB
    ws.row_dimensions[3].height = 14

    # ══ 1) DETALLE POR PRODUCTO ══
    row = _section(ws, 5, "DETALLE POR PRODUCTO (USD)", nd)
    hrow = row
    for j, (name, width, _kind) in enumerate(DETALLE_COLS, start=1):
        c = ws.cell(row=row, column=j, value=name)
        c.font = FONT_HEAD
        c.fill = FILL_NEGRO
        c.alignment = CENTER
        c.border = BORDER_ALL
        ws.column_dimensions[get_column_letter(j)].width = width
    ws.row_dimensions[row].height = 30
    row += 1
    first_data = row
    for i, it in enumerate(calc["items"], start=1):
        esp = bool(it.get("especial"))
        vals = [
            i, it["nombre"], float(it["cant"]), float(it["valor_unit"]), float(it["fob"]),
            float(it["ratio"]) / 100,
            float(it["advalorem_pct_eff"]) / 100, float(it["seguro_pct_eff"]) / 100,
            float(it["fondinfa_pct_eff"]) / 100, float(it["iva_pct_eff"]) / 100,
            float(it["advalorem"]), float(it["flete"]), float(it["seguro"]),
            float(it["fondinfa"]), float(it["adicionales"]), float(it["valor_total"]),
            float(it["costo_unit"]),
            float(it["compet1"]) / 100, float(it["comp1_sin"]), float(it["comp1_con"]),
            float(it["compet2"]) / 100, float(it["comp2_sin"]), float(it["comp2_con"]),
            float(it["compet3"]) / 100, float(it["comp3_sin"]), float(it["comp3_con"]),
            "Especial" if esp else "General",
        ]
        banda = FILL_BANDA if (i % 2 == 0 and not esp) else None
        for j, ((_name, _w, kind), v) in enumerate(zip(DETALLE_COLS, vals), start=1):
            c = ws.cell(row=row, column=j, value=v)
            c.border = BORDER_ALL
            if esp:
                c.fill = FILL_AMBAR
            elif banda:
                c.fill = banda
            if kind == "money":
                c.number_format = MONEY_FMT
                c.alignment = RIGHT
                c.font = FONT_BODY
            elif kind == "pct":
                c.number_format = PCT_FMT
                c.alignment = RIGHT
                c.font = FONT_BODY
            elif kind == "num":
                c.number_format = '#,##0.00'
                c.alignment = RIGHT
                c.font = FONT_BODY
            elif kind == "left":
                c.alignment = LEFT
                c.font = FONT_BODY
            elif kind == "center":
                c.alignment = CENTER
                c.font = Font(name="Calibri", size=9, bold=True,
                              color="92400E" if esp else "52525B")
            else:
                c.alignment = CENTER
                c.font = FONT_BODY
        for j in (19, 20, 22, 23, 25, 26):
            ws.cell(row=row, column=j).font = FONT_VERDE
        ws.row_dimensions[row].height = 16
        row += 1
    # Fila de totales del detalle (negro + dorado)
    tot_row = row
    c0 = ws.cell(row=row, column=1, value="TOTAL")
    c0.font = FONT_DORADO
    c0.fill = FILL_NEGRO
    c0.border = BORDER_ALL
    c0.alignment = CENTER
    for j in range(2, nd + 1):
        cc = ws.cell(row=row, column=j)
        cc.fill = FILL_NEGRO
        cc.border = BORDER_ALL
    col_letter = {j: get_column_letter(j) for j in range(1, nd + 1)}
    for j in (5, 11, 12, 13, 14, 15, 16, 17):
        L = col_letter[j]
        cc = ws.cell(row=row, column=j)
        cc.value = "=SUM(%s%d:%s%d)" % (L, first_data, L, row - 1)
        cc.number_format = MONEY_FMT
        cc.font = FONT_DORADO
        cc.alignment = RIGHT
    ws.row_dimensions[row].height = 18
    ws.auto_filter.ref = "%s%d:%s%d" % (
        col_letter[1], hrow, col_letter[nd], tot_row - 1)
    ws.freeze_panes = "C%d" % first_data
    _marco(ws, hrow, tot_row, 1, nd)
    row += 2

    # ══ 2) RESUMEN GENERAL ══
    ncols = 3
    ws.column_dimensions["A"].width = 34
    ws.column_dimensions["B"].width = 18
    ws.column_dimensions["C"].width = 30
    sec_row = row
    row = _section(ws, row, "RESUMEN GENERAL (USD)", ncols)

    row = _subsec(ws, row, "Tributos de importación", ncols)
    tribs = [
        ("FOB", R["fob"], False, None),
        ("Flete internacional", R["flete_intl"], False, None),
        ("Seguro (%.2f %%)" % T.get("seguro_pct", 0), R["seguro"], False, None),
        ("Ad-Valorem (%.2f %%)" % T.get("advalorem_pct", 0), R["advalorem"], False, None),
        ("Fondinfa (%.2f %%)" % T.get("fondinfa_pct", 0), R["fondinfa"], False, None),
        ("IVA importación (%.2f %%)" % T.get("iva_pct", 0), R["iva"], False, None),
        ("Total tributos", R["total_tributos"], True, FILL_GRIS),
    ]
    for label, value, bold, fill in tribs:
        row = _kv_row(ws, row, label, float(value), bold, fill, ncols)

    row = _subsec(ws, row, "Gastos locales · origen + destino", ncols)
    locales = [
        ("ISD / Comisión banco", R["comision_banco"], False, None),
        ("Flete local", R["flete_local"], False, None),
        ("Bodegaje y manipulación", R["bodegaje"], False, None),
        ("Devolución contenedor", R["devolucion"], False, None),
        ("Gastos administrativos", R["gasto_admin"], False, None),
        ("Descarga / almuerzo", R["descarga"], False, None),
        ("Agente aduana", R["agente"], False, None),
        ("Otros", R["otros"], False, None),
        ("Gastos locales", R["locales_total"], True, FILL_GRIS),
    ]
    for label, value, bold, fill in locales:
        row = _kv_row(ws, row, label, float(value), bold, fill, ncols)

    gran_row = row
    row = _kv_row(ws, row, "GRAN TOTAL GN+GL", float(R["gran_total"]),
                  True, FILL_NEGRO, ncols)
    for col in (1, 2):
        ws.cell(row=gran_row, column=col).font = FONT_DORADO
    row = _kv_row(ws, row, "Impuestos + IVA", float(R["impuesto_iva_sum"]),
                  True, FILL_AMBAR, ncols)
    _marco(ws, sec_row, row - 1, 1, ncols)

    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()
