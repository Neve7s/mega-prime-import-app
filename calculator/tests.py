from decimal import Decimal

from django.test import TestCase

from .models import HSCode, TaxRule
from .reports import build_reporte_xlsx
from .seed_embarques import EMBARQUES
from .services import calculate_shipment, estimate


class EstimateTest(TestCase):
    def test_caso_manual(self):
        # FOB 10.000 + flete 1.500 + seguro 100 = CIF 11.600
        # ad-valorem 10% = 1.160 | fodinfa 0.5% = 58 | base IVA 12.818 | IVA 15% = 1.922,70
        r = estimate(Decimal("10000"), Decimal("1500"), Decimal("100"), Decimal("10"))
        self.assertEqual(r["cif"], Decimal("11600.00"))
        self.assertEqual(r["ad_valorem"], Decimal("1160.00"))
        self.assertEqual(r["fodinfa"], Decimal("58.00"))
        self.assertEqual(r["iva"], Decimal("1922.70"))
        self.assertEqual(r["total"], Decimal("14740.70"))

    def test_regla_activa(self):
        TaxRule.objects.create(nombre="Régimen general 2026")
        rule = TaxRule.get_active()
        self.assertIsNotNone(rule)
        self.assertEqual(rule.iva_pct, Decimal("15.00"))
        self.assertTrue(HSCode.objects.create(codigo="6907.21", descripcion="Placas cerámicas", ad_valorem_pct=Decimal("15")))


class ShipmentExcelTest(TestCase):
    """Fixtures reales de docs/01-analisis-excel-costos.md (Teja Roma)."""

    def test_replica_teja_roma(self):
        items = [
            {"nombre": "ROMA 7.216", "cant": 600, "cant_caja": 1, "valor_unit": "20.31", "margen_pct": 35},
            {"nombre": "ROMA 6.232", "cant": 675, "cant_caja": 1, "valor_unit": "17.54", "margen_pct": 35},
            {"nombre": "ROMA 5.248", "cant": 735, "cant_caja": 1, "valor_unit": "14.77", "margen_pct": 35},
            {"nombre": "MAIN RIDGE", "cant": 90, "cant_caja": 10, "valor_unit": "27.5", "margen_pct": 40},
            {"nombre": "TORNILLOS", "cant": 15, "cant_caja": 3000, "valor_unit": 90, "margen_pct": 40},
        ]
        # Subtotales AH del Excel (AJ/1.15 donde aplica IVA)
        gastos = {
            "flete_intl": "4775", "flete_local": "4650",
            "gasto_admin": Decimal("2407.816") / Decimal("1.15"),
            "bodegaje": Decimal("1379.766") / Decimal("1.15"),
            "devolucion": Decimal("330.75") / Decimal("1.15"),
            "comision_banco": "4597.666", "descarga": "1483.333",
            "agente": Decimal("207.573") / Decimal("1.15"), "otros_unit": 0,
        }
        tasas = {"seguro_pct": 1, "advalorem_pct": 20, "fondinfa_pct": "0.5",
                 "iva_pct": 15, "iva_pvp_pct": 15}
        r = calculate_shipment(items, gastos, tasas)
        self.assertIsNone(r["error"])
        # Proporciones verificadas contra Excel con estos 5 items parciales:
        # ratio del primero = 600/total
        total_cant = 600 + 675 + 735 + 90 + 15
        self.assertEqual(r["resumen"]["total_cant"], Decimal(str(total_cant) + ".00"))
        # Suma de ratios = 100%
        suma_ratio = sum(i["ratio"] for i in r["items"])
        self.assertTrue(Decimal("99.99") <= suma_ratio <= Decimal("100.01"))
        # El 35% heredado se interpreta como recargo sobre el costo unitario.
        primero = r["items"][0]
        esperado = (primero["costo_unit"] * Decimal("1.35")
                    * Decimal("1.15")).quantize(Decimal("0.01"))
        self.assertEqual(primero["comp1_con"], esperado)


class ShipmentOverrideTest(TestCase):
    """Override de impuestos por producto: vacío = tasa general."""

    def test_fob_es_cantidad_por_precio_unitario(self):
        result = calculate_shipment(
            [{"nombre": "A", "cant": 3, "valor_unit": "12.50"}],
            {"flete_intl": 0}, {"seguro_pct": 0, "advalorem_pct": 0,
                                "fondinfa_pct": 0, "iva_pct": 0})
        self.assertEqual(result["resumen"]["fob"], Decimal("37.50"))

    def test_recargos_por_cliente_son_porcentaje_sobre_costo(self):
        result = calculate_shipment(
            [{"nombre": "A", "cant": 2, "valor_unit": "10",
              "compet1": 40, "compet2": 20, "compet3": 0}],
            {}, {"seguro_pct": 0, "advalorem_pct": 0, "fondinfa_pct": 0,
                 "iva_pct": 0, "iva_pvp_pct": 15})
        item = result["items"][0]
        self.assertEqual(item["costo_unit"], Decimal("10.00"))
        self.assertEqual(item["comp1_sin"], Decimal("14.00"))
        self.assertEqual(item["comp1_con"], Decimal("16.10"))
        self.assertEqual(item["comp2_sin"], Decimal("12.00"))
        self.assertEqual(item["comp2_con"], Decimal("13.80"))
        self.assertEqual(item["comp3_sin"], Decimal("10.00"))
        self.assertEqual(item["comp3_con"], Decimal("11.50"))

    def test_costo_unitario_incluye_iva_pagado_en_importacion(self):
        result = calculate_shipment(
            [{"nombre": "A", "cant": 2, "valor_unit": "10", "compet1": 40}],
            {}, {"seguro_pct": 0, "advalorem_pct": 0, "fondinfa_pct": 0,
                 "iva_pct": 15, "iva_pvp_pct": 15})
        item = result["items"][0]
        self.assertEqual(item["valor_total"], Decimal("23.00"))
        self.assertEqual(item["costo_unit"], Decimal("11.50"))
        self.assertEqual(item["comp1_sin"], Decimal("16.10"))
        self.assertEqual(item["comp1_con"], Decimal("18.52"))

    def _base(self):
        items = [
            {"nombre": "A", "cant": 100, "cant_caja": 1, "valor_unit": "10", "margen_pct": 35},
            {"nombre": "B", "cant": 100, "cant_caja": 1, "valor_unit": "10", "margen_pct": 35},
        ]
        gastos = {"flete_intl": "1000", "flete_local": 0, "gasto_admin": 0,
                  "bodegaje": 0, "devolucion": 0, "comision_banco": 0,
                  "descarga": 0, "agente": 0, "otros_unit": 0}
        tasas = {"seguro_pct": 1, "advalorem_pct": 20, "fondinfa_pct": "0.5",
                 "iva_pct": 15, "iva_pvp_pct": 15}
        return items, gastos, tasas

    def test_sin_override_igual_reparto(self):
        items, gastos, tasas = self._base()
        r = calculate_shipment(items, gastos, tasas)
        self.assertIsNone(r["error"])
        a, b = r["items"]
        self.assertFalse(a["especial"])
        self.assertEqual(a["advalorem"], b["advalorem"])
        self.assertEqual(a["advalorem_pct_eff"], Decimal("20.00"))

    def test_override_advalorem_solo_un_producto(self):
        items, gastos, tasas = self._base()
        items[1]["advalorem_pct"] = "0"  # B exonerado, A sigue al 20%
        r = calculate_shipment(items, gastos, tasas)
        a, b = r["items"]
        self.assertFalse(a["especial"])
        self.assertTrue(b["especial"])
        self.assertEqual(b["advalorem"], Decimal("0.00"))
        # A = (fob_A + seguro_A) * 20%. fob_A=1000, flete_A=500, seguro_A=15
        self.assertEqual(a["advalorem"], Decimal("203.00"))
        # Total = solo el de A
        self.assertEqual(r["resumen"]["advalorem"], Decimal("203.00"))
        self.assertEqual(b["advalorem_pct_eff"], Decimal("0.00"))

    def test_semilla_excel_cuadra_totales(self):
        for key, esperado in (("roma", Decimal("99922.97")),
                              ("espanol", Decimal("133689.92"))):
            e = EMBARQUES[key]
            r = calculate_shipment(e["items"], e["gastos"], e["tasas"])
            self.assertIsNone(r["error"])
            self.assertEqual(r["resumen"]["gran_total"], esperado)

    def test_override_vacio_usa_general(self):
        items, gastos, tasas = self._base()
        items[0]["advalorem_pct"] = ""
        items[1]["seguro_pct"] = None
        r = calculate_shipment(items, gastos, tasas)
        self.assertFalse(r["items"][0]["especial"])
        self.assertFalse(r["items"][1]["especial"])
        self.assertEqual(r["items"][0]["advalorem"], r["items"][1]["advalorem"])


class ReporteXlsxTest(TestCase):
    def test_endpoint_reporte_exporta_los_tres_precios_por_cliente(self):
        import json
        from io import BytesIO

        from django.contrib.auth.models import User
        from openpyxl import load_workbook

        user = User.objects.create_user(
            "report-margins", "report-margins-password", is_staff=True)
        self.client.force_login(user)
        payload = {
            "embarque": "Márgenes",
            "tasas": {"seguro_pct": 0, "advalorem_pct": 0, "fondinfa_pct": 0,
                      "iva_pct": 0, "iva_pvp_pct": 15},
            "gastos": {},
            "items": [{"nombre": "Producto", "cant": 1, "valor_unit": 10,
                       "compet1": 40, "compet2": 20, "compet3": 0}],
        }
        response = self.client.post(
            "/webadmin/calculadora/reporte/", json.dumps(payload),
            content_type="application/json")
        self.assertEqual(response.status_code, 200)
        ws = load_workbook(BytesIO(response.content))["Reporte"]
        header_row = next(
            row for row in ws.iter_rows()
            if "PVP May. s/IGV" in [cell.value for cell in row])
        headers = [cell.value for cell in header_row]
        values = [cell.value for cell in ws[header_row[0].row + 1]]
        self.assertEqual(values[headers.index("PVP May. s/IGV")], 14)
        self.assertEqual(values[headers.index("PVP May. c/IGV")], 16.1)
        self.assertEqual(values[headers.index("PVP Frec. s/IGV")], 12)
        self.assertEqual(values[headers.index("PVP Event. s/IGV")], 10)

    def test_reporte_con_estilos_y_totales(self):
        import json
        from io import BytesIO

        from openpyxl import load_workbook

        e = EMBARQUES["roma"]
        calc = calculate_shipment(e["items"], e["gastos"], e["tasas"])
        data = build_reporte_xlsx(calc, e["tasas"], e["gastos"], "Teja Roma")
        wb = load_workbook(BytesIO(data))
        ws = wb["Reporte"]
        self.assertEqual(ws["A1"].value, "MEGA PRIME IMPORT")
        self.assertEqual(ws["A1"].fill.fgColor.rgb[-6:], "0A0A0B")
        gran = None
        for row in ws.iter_rows(min_row=1, max_row=100, max_col=2):
            if row[0].value == "GRAN TOTAL GN+GL":
                gran = row[1].value
        self.assertAlmostEqual(gran, 99922.97, places=2)
        # Encabezado del detalle con fondo negro corporativo
        found = False
        for row in ws.iter_rows():
            vals = [c.value for c in row]
            if "PVP May. c/IGV" in vals:
                found = True
                self.assertEqual(row[0].fill.fgColor.rgb[-6:], "0A0A0B")
                break
        self.assertTrue(found)

    def test_endpoint_descarga_xlsx(self):
        import json
        from io import BytesIO

        from django.contrib.auth.models import User
        from openpyxl import load_workbook

        User.objects.create_user("rep", password="x", is_staff=True)
        self.client.force_login(User.objects.get(username="rep"))
        e = EMBARQUES["espanol"]
        payload = {"embarque": "Teja Español", "tasas": e["tasas"],
                   "gastos": e["gastos"], "items": e["items"]}
        r = self.client.post("/webadmin/calculadora/reporte/", json.dumps(payload),
                             content_type="application/json")
        self.assertEqual(r.status_code, 200)
        self.assertIn("spreadsheetml.sheet", r["Content-Type"])
        self.assertIn(".xlsx", r["Content-Disposition"])
        ws = load_workbook(BytesIO(r.content))["Reporte"]
        gran = next(row[1].value for row in
                    ws.iter_rows(min_row=1, max_row=100, max_col=2)
                    if row[0].value == "GRAN TOTAL GN+GL")
        self.assertAlmostEqual(gran, 133689.92, places=2)


class PlantillaImportacionTest(TestCase):
    def test_plantilla_tiene_solo_gastos_y_productos(self):
        from io import BytesIO

        from openpyxl import load_workbook

        from .excel_templates import build_plantilla_xlsx

        workbook = load_workbook(BytesIO(build_plantilla_xlsx()))
        self.assertEqual(workbook.sheetnames, ["Gastos", "Productos"])
        self.assertEqual(
            [cell.value for cell in workbook["Productos"][1]],
            ["Producto", "Cantidad", "Precio FOB unitario USD"])
        self.assertEqual(
            [cell.value for cell in workbook["Gastos"][1]],
            ["Servicio", "Valor USD"])

    def test_importa_gastos_cantidad_y_precio(self):
        from io import BytesIO

        from openpyxl import load_workbook

        from .excel_templates import build_plantilla_xlsx, read_plantilla_xlsx

        workbook = load_workbook(BytesIO(build_plantilla_xlsx()))
        workbook["Gastos"]["B2"] = 1200
        workbook["Gastos"]["B4"] = 25
        workbook["Productos"]["A2"] = "Piedra"
        workbook["Productos"]["B2"] = 10
        workbook["Productos"]["C2"] = 8.5
        output = BytesIO()
        workbook.save(output)

        imported = read_plantilla_xlsx(BytesIO(output.getvalue()))
        self.assertEqual(imported["gastos"]["flete_intl"], "1200")
        self.assertEqual(imported["gastos"]["comision_banco"], "25")
        self.assertEqual(imported["items"][0]["cant"], "10")
        self.assertEqual(imported["items"][0]["valor_unit"], "8.5")
        self.assertEqual(imported["items"][0]["compet1"], "")
        self.assertEqual(imported["items"][0]["compet2"], "")
        self.assertEqual(imported["items"][0]["compet3"], "")


class SavedShipmentViewTest(TestCase):
    def setUp(self):
        from django.contrib.auth.models import User

        self.user = User.objects.create_user(
            username="shipment-staff", password="test-password", is_staff=True)
        self.client.force_login(self.user)
        self.tasas = {
            "seguro_pct": "1", "advalorem_pct": "20", "fondinfa_pct": "0.5",
            "iva_pct": "15", "iva_pvp_pct": "15",
        }
        self.gastos = {
            "flete_intl": "0", "flete_local": "0", "gasto_admin": "0",
            "bodegaje": "0", "devolucion": "0", "comision_banco": "0",
            "descarga": "0", "agente": "0", "otros_unit": "0",
        }

    def test_guarda_y_carga_embarque_vacio(self):
        import json

        from .models import SavedShipment

        payload = {
            "nombre": "Compra de prueba", "embarque_key": "nuevo",
            "tasas": self.tasas, "gastos": self.gastos,
            "items": [{"nombre": "", "cant": "1", "cant_caja": "1",
                       "valor_unit": "", "margen_pct": "35"}],
        }
        response = self.client.post(
            "/webadmin/calculadora/guardar/", json.dumps(payload),
            content_type="application/json")
        self.assertEqual(response.status_code, 200)
        saved = SavedShipment.objects.get()
        self.assertEqual(saved.nombre, "Compra de prueba")
        self.assertEqual(response.json()["key"], f"guardado-{saved.pk}")

        page = self.client.get(f"/webadmin/calculadora/?embarque=guardado-{saved.pk}")
        self.assertEqual(page.status_code, 200)
        self.assertContains(page, 'value="Compra de prueba"')
        self.assertContains(page, "May %")
        self.assertContains(page, "May $")
        self.assertContains(page, "Frec %")
        self.assertContains(page, "Frec $")
        self.assertContains(page, "Eve %")
        self.assertContains(page, "PVP $")
        self.assertNotContains(page, "C/Caja")

        payload["embarque_key"] = f"guardado-{saved.pk}"
        payload["nombre"] = "Compra actualizada"
        response = self.client.post(
            "/webadmin/calculadora/guardar/", json.dumps(payload),
            content_type="application/json")
        self.assertEqual(response.status_code, 200)
        self.assertEqual(SavedShipment.objects.count(), 1)
        saved.refresh_from_db()
        self.assertEqual(saved.nombre, "Compra actualizada")

    def test_plantilla_descarga_y_subida(self):
        from io import BytesIO

        from django.core.files.uploadedfile import SimpleUploadedFile
        from openpyxl import load_workbook

        from .excel_templates import build_plantilla_xlsx

        response = self.client.get("/webadmin/calculadora/plantilla/")
        self.assertEqual(response.status_code, 200)
        self.assertIn("spreadsheetml.sheet", response["Content-Type"])

        workbook = load_workbook(BytesIO(build_plantilla_xlsx()))
        workbook["Gastos"]["B2"] = 750
        workbook["Productos"]["A2"] = "Piedra"
        workbook["Productos"]["B2"] = 4
        workbook["Productos"]["C2"] = 12
        workbook_data = BytesIO()
        workbook.save(workbook_data)
        response = self.client.post(
            "/webadmin/calculadora/importar/",
            {"archivo": SimpleUploadedFile(
                "plantilla.xlsx", workbook_data.getvalue(),
                content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
        )
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.json()["items"][0]["cant"], "4")
        self.assertEqual(response.json()["items"][0]["valor_unit"], "12")
        self.assertEqual(response.json()["gastos"]["flete_intl"], "750")


class EmbarqueGestiónTest(TestCase):
    """Renombrar y eliminar embarques guardados."""

    def setUp(self):
        import json

        from django.contrib.auth.models import User

        from .models import SavedShipment

        self.user = User.objects.create_user(
            username="gestion-staff", password="test-password", is_staff=True)
        self.client.force_login(self.user)
        self.shipment = SavedShipment.objects.create(
            nombre="Compra vieja", tasas={}, gastos={}, items=[])
        self.key = f"guardado-{self.shipment.pk}"
        self.payload = json.dumps({"embarque_key": self.key, "nombre": "Compra nueva"})

    def test_renombrar_embarque_guardado(self):
        response = self.client.post(
            "/webadmin/calculadora/renombrar/", self.payload,
            content_type="application/json")
        self.assertEqual(response.status_code, 200)
        self.shipment.refresh_from_db()
        self.assertEqual(self.shipment.nombre, "Compra nueva")
        self.assertEqual(response.json()["key"], self.key)

    def test_renombrar_exige_nombre(self):
        import json

        response = self.client.post(
            "/webadmin/calculadora/renombrar/",
            json.dumps({"embarque_key": self.key, "nombre": "   "}),
            content_type="application/json")
        self.assertEqual(response.status_code, 400)
        self.shipment.refresh_from_db()
        self.assertEqual(self.shipment.nombre, "Compra vieja")

    def test_no_se_renombra_una_semilla(self):
        import json

        response = self.client.post(
            "/webadmin/calculadora/renombrar/",
            json.dumps({"embarque_key": "roma", "nombre": "Otra"}),
            content_type="application/json")
        self.assertEqual(response.status_code, 400)

    def test_eliminar_embarque_guardado(self):
        import json

        from .models import SavedShipment

        response = self.client.post(
            "/webadmin/calculadora/eliminar/", json.dumps({"embarque_key": self.key}),
            content_type="application/json")
        self.assertEqual(response.status_code, 200)
        self.assertEqual(SavedShipment.objects.count(), 0)

    def test_no_se_elimina_una_semilla(self):
        import json

        from .models import SavedShipment

        response = self.client.post(
            "/webadmin/calculadora/eliminar/", json.dumps({"embarque_key": "roma"}),
            content_type="application/json")
        self.assertEqual(response.status_code, 400)
        self.assertEqual(SavedShipment.objects.count(), 1)

    def test_embarque_inexistente_devuelve_json(self):
        import json

        from .models import SavedShipment

        response = self.client.post(
            "/webadmin/calculadora/eliminar/",
            json.dumps({"embarque_key": "guardado-999999"}),
            content_type="application/json")
        self.assertEqual(response.status_code, 400)
        self.assertIn("error", response.json())
        self.assertEqual(SavedShipment.objects.count(), 1)

    def test_clave_con_digitos_no_ascii_no_revienta(self):
        import json

        response = self.client.post(
            "/webadmin/calculadora/eliminar/",
            json.dumps({"embarque_key": "guardado-²"}),
            content_type="application/json")
        self.assertEqual(response.status_code, 400)

    def test_la_calculadora_tiene_busqueda_y_botones(self):
        page = self.client.get("/webadmin/calculadora/").content.decode("utf-8")
        self.assertIn('id="embarque-buscar"', page)
        self.assertIn('id="btn-renombrar"', page)
        self.assertIn('id="btn-eliminar"', page)
