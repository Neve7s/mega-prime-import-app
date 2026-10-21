from django.test import TestCase


class PanelTest(TestCase):
    def setUp(self):
        from django.contrib.auth.models import User

        self.staff = User.objects.create_user("webadmin", password="x", is_staff=True)
        self.plain = User.objects.create_user("nadie", password="x")

    def test_anonimo_va_a_login(self):
        resp = self.client.get("/webadmin/")
        self.assertEqual(resp.status_code, 302)
        self.assertIn("/webadmin/login/", resp["Location"])

    def test_lista_productos_requiere_staff(self):
        resp = self.client.get("/webadmin/productos/")
        self.assertEqual(resp.status_code, 302)
        self.assertIn("/webadmin/login/", resp["Location"])
        self.client.force_login(self.plain)
        resp = self.client.get("/webadmin/productos/")
        self.assertEqual(resp.status_code, 302)
        self.assertIn("/webadmin/login/", resp["Location"])

    def test_sin_staff_no_entra(self):
        self.client.force_login(self.plain)
        resp = self.client.get("/webadmin/")
        self.assertEqual(resp.status_code, 302)

    def test_staff_ve_tablero_y_productos(self):
        self.client.force_login(self.staff)
        self.assertEqual(self.client.get("/webadmin/").status_code, 200)
        self.assertEqual(self.client.get("/webadmin/productos/").status_code, 200)
        self.assertEqual(self.client.get("/webadmin/pedidos/").status_code, 200)
        self.assertEqual(self.client.get("/webadmin/codigos/").status_code, 200)
        self.assertEqual(self.client.get("/webadmin/ajustes/").status_code, 200)
        self.assertEqual(self.client.get("/webadmin/calculadora/").status_code, 200)

    def test_ajustes_subtabs_y_guardado(self):
        self.client.force_login(self.staff)
        html = self.client.get("/webadmin/ajustes/").content.decode()
        for tab in ("marca", "portada", "contacto", "redes", "seo", "pie", "tema"):
            self.assertIn('data-tab="%s"' % tab, html)
        self.assertIn("Paletas predefinidas", html)
        self.assertIn('data-drop', html)  # dropzones de imagen
        resp = self.client.post("/webadmin/ajustes/", data={
            "site_name": "MP", "meta_title": "T", "tagline": "S",
            "company_legal": "E", "hero_city": "C", "hero_title": "HT",
            "hero_cta_primary": "A", "hero_cta_secondary": "B",
            "whatsapp_number": "593900000000", "phone_display": "+593",
            "email": "v@v.com", "address": "D", "cities": "C",
            "schedule": "H", "footer_text": "F",
            "color_primario": "#D4AF37", "color_secundario": "#0A0A0B",
            "color_fondo": "#FAFAF8", "color_tarjetas": "#FFFFFF",
            "color_texto": "#131313", "color_acento": "#15803D",
            "fuente_cuerpo": "Inter", "fuente_titulos": "Montserrat",
        })
        self.assertEqual(resp.status_code, 200)
        self.assertIn("Guardado", resp.content.decode())
        from core.models import SiteSettings
        site = SiteSettings.get_solo()
        self.assertEqual(site.fuente_titulos, "Montserrat")
        self.assertEqual(site.color_primario, "#D4AF37")

    def test_ajustes_sube_logo(self):
        import io
        import shutil
        import tempfile

        from django.core.files.uploadedfile import SimpleUploadedFile
        from django.test import override_settings
        from PIL import Image

        tmp = tempfile.mkdtemp()
        try:
            with override_settings(MEDIA_ROOT=tmp):
                buf = io.BytesIO()
                Image.new("RGB", (100, 60), (212, 175, 55)).save(buf, "PNG")
                logo = SimpleUploadedFile("logo-test.png", buf.getvalue(), "image/png")
                self.client.force_login(self.staff)
                resp = self.client.post("/webadmin/ajustes/", data={
                    "site_name": "MP", "meta_title": "T", "tagline": "S",
                    "company_legal": "E", "hero_city": "C", "hero_title": "HT",
                    "hero_cta_primary": "A", "hero_cta_secondary": "B",
                    "whatsapp_number": "593900000000", "phone_display": "+593",
                    "email": "v@v.com", "address": "D", "cities": "C",
                    "schedule": "H", "footer_text": "F",
                    "color_primario": "#D4AF37", "color_secundario": "#0A0A0B",
                    "color_fondo": "#FAFAF8", "color_tarjetas": "#FFFFFF",
                    "color_texto": "#131313", "color_acento": "#15803D",
                    "fuente_cuerpo": "Inter", "fuente_titulos": "Inter",
                    "logo": logo,
                })
                self.assertEqual(resp.status_code, 200)
                from core.models import SiteSettings
                site = SiteSettings.get_solo()
                self.assertTrue(bool(site.logo))
                self.assertIn("logo-test", site.logo.name)
        finally:
            shutil.rmtree(tmp, ignore_errors=True)

    def test_categorias_crud_y_filtro_automatico(self):
        from catalog.models import Category, Product

        self.client.force_login(self.staff)
        # Crear
        r = self.client.post("/webadmin/categorias/nueva/", data={
            "nombre": "Grifería", "slug": "griferia",
            "descripcion": "", "orden": "3",
        })
        self.assertEqual(r.status_code, 302)
        cat = Category.objects.get(slug="griferia")
        # Aparece sola en el filtro del catálogo público (con productos)
        producto = Product.objects.create(categoria=cat, sku="GR-T", nombre="Grifo test",
                                          slug="grifo-test", precio_mayor="10.00")
        html = self.client.get("/catalogo/").content.decode()
        self.assertIn('value="griferia"', html)
        self.assertIn("Grifería", html)
        # Y en el selector del formulario de producto
        html_p = self.client.get("/webadmin/productos/nuevo/").content.decode()
        self.assertIn("Grifería", html_p)
        # Editar
        r = self.client.post(f"/webadmin/categorias/{cat.pk}/", data={
            "nombre": "Griferías Finas", "slug": "griferia",
            "descripcion": "", "orden": "1",
        })
        self.assertEqual(r.status_code, 302)
        cat.refresh_from_db()
        self.assertEqual(cat.nombre, "Griferías Finas")
        # Eliminar vacía
        producto.delete()
        r = self.client.post(f"/webadmin/categorias/{cat.pk}/eliminar/")
        self.assertEqual(r.status_code, 302)
        self.assertFalse(Category.objects.filter(slug="griferia").exists())

    def test_el_tipo_define_la_categoria(self):
        """Una sola decisión: al elegir el tipo, la categoría se pone sola."""
        from catalog.models import Category, Product, TipoProducto

        self.client.force_login(self.staff)
        tejas = Category.objects.create(nombre="Tejas", slug="tejas-td")
        tipo = TipoProducto.objects.create(nombre="Teja Española", slug="teja-esp-td",
                                           categoria=tejas)
        r = self.client.post("/webadmin/productos/nuevo/", data={
            "tipo": tipo.pk, "nombre": "Teja sin categoría", "descripcion": "",
            "precio_mayor": "12.50", "unidad": "m2", "stock_cantidad": "10",
            "en_stock": "on", "activo": "on",
            "colores_json": "[]", "caracteristicas_json": "[]",
        })
        self.assertEqual(r.status_code, 302)
        p = Product.objects.get(nombre="Teja sin categoría")
        self.assertEqual(p.categoria, tejas)
        # El SKU y el slug se generan solos
        self.assertTrue(p.sku.startswith("COD-"))
        self.assertEqual(p.slug, "teja-sin-categoria")

    def test_sin_tipo_hace_falta_la_categoria(self):
        from catalog.models import Category

        self.client.force_login(self.staff)
        Category.objects.create(nombre="Tejas", slug="tejas-sc")
        r = self.client.post("/webadmin/productos/nuevo/", data={
            "nombre": "Producto sin nada", "descripcion": "",
            "precio_mayor": "10", "unidad": "m2", "stock_cantidad": "0",
            "activo": "on", "colores_json": "[]", "caracteristicas_json": "[]",
        })
        self.assertEqual(r.status_code, 200)
        self.assertContains(r, "Elige el tipo de producto o la categoría.")

    def test_producto_guarda_colores_y_caracteristicas(self):
        import json

        from catalog.models import Category, Color, Product

        self.client.force_login(self.staff)
        cat = Category.objects.create(nombre="Tejas", slug="tejas-t")
        rojo = Color.objects.get(slug="rojo")  # viene de la paleta
        base = {"categoria": cat.pk, "sku": "MPG-COL", "nombre": "Con color",
                "slug": "con-color", "descripcion": "", "precio_mayor": "10.00",
                "unidad": "m2", "stock_cantidad": "1", "imagen_url": "",
                "en_stock": "on", "activo": "on",
                "colores_json": json.dumps([
                    {"slug": "rojo", "nombre": "Rojo", "hex": "#8B2500", "rol": "exterior"},
                    {"slug": "", "nombre": "Malo", "hex": "zzz", "rol": ""},
                    {"slug": "", "nombre": "", "hex": "#FFFFFF", "rol": ""}]),
                "caracteristicas_json": json.dumps(
                    ["Resistente UV", "  ", "Peso ligero"])}
        r = self.client.post("/webadmin/productos/nuevo/", data=base)
        self.assertEqual(r.status_code, 302)
        p = Product.objects.get(sku="MPG-COL")
        colores = list(p.colores_producto.all())
        self.assertEqual(len(colores), 1)
        self.assertEqual(colores[0].color, rojo)
        self.assertEqual(colores[0].rol, "exterior")
        self.assertEqual(p.caracteristicas, ["Resistente UV", "Peso ligero"])
        det = self.client.get("/catalogo/producto/con-color/").content.decode()
        self.assertIn("Rojo", det)
        self.assertIn("Resistente UV", det)
        self.assertIn(rojo.hex, det)

    def test_producto_guarda_galeria_por_url_y_archivo(self):
        import shutil
        import tempfile
        from io import BytesIO

        from django.core.files.uploadedfile import SimpleUploadedFile
        from django.test import override_settings
        from PIL import Image

        from catalog.models import Category, Product, ProductImage

        self.client.force_login(self.staff)
        cat = Category.objects.create(nombre="Tejas", slug="tejas-gal")
        buffer = BytesIO()
        Image.new("RGB", (40, 40), "red").save(buffer, "PNG")
        base = {"categoria": cat.pk, "sku": "MPG-GAL", "nombre": "Con galería",
                "slug": "con-galeria", "descripcion": "", "precio_mayor": "10.00",
                "unidad": "m2", "stock_cantidad": "1", "imagen_url": "",
                "en_stock": "on", "activo": "on",
                "colores_json": "[]", "caracteristicas_json": "[]",
                "galeria_urls": "https://ejemplo.com/foto1.jpg\nno-es-url\nhttps://ejemplo.com/foto2.jpg",
                "galeria_archivos": SimpleUploadedFile(
                    "foto.png", buffer.getvalue(), content_type="image/png")}
        destino = tempfile.mkdtemp()
        try:
            with override_settings(MEDIA_ROOT=destino):
                r = self.client.post("/webadmin/productos/nuevo/", data=base)
                self.assertEqual(r.status_code, 302)
                p = Product.objects.get(sku="MPG-GAL")
                fotos = list(p.imagenes.all())
                self.assertEqual(len(fotos), 3)
                self.assertEqual(fotos[0].imagen_url, "https://ejemplo.com/foto1.jpg")
                self.assertEqual(fotos[1].imagen_url, "https://ejemplo.com/foto2.jpg")
                self.assertTrue(fotos[2].imagen)
                # La galería se ve en la ficha pública
                det = self.client.get("/catalogo/producto/con-galeria/").content.decode()
                self.assertIn("galeria-mini", det)
                self.assertIn("https://ejemplo.com/foto1.jpg", det)
        finally:
            shutil.rmtree(destino, ignore_errors=True)

    def test_producto_borra_fotos_marcadas(self):
        from catalog.models import Category, Product, ProductImage

        self.client.force_login(self.staff)
        cat = Category.objects.create(nombre="Tejas", slug="tejas-bor")
        p = Product.objects.create(categoria=cat, sku="MPG-BOR", nombre="A borrar",
                                   slug="a-borrar", precio_mayor="10")
        foto = ProductImage.objects.create(producto=p, imagen_url="https://ejemplo.com/x.jpg")
        base = {"categoria": cat.pk, "sku": p.sku, "nombre": p.nombre, "slug": p.slug,
                "descripcion": "", "precio_mayor": "10.00", "unidad": "m2",
                "stock_cantidad": "0", "imagen_url": "", "activo": "on",
                "colores_json": "[]", "caracteristicas_json": "[]",
                "galeria_borrar": str(foto.pk)}
        r = self.client.post(f"/webadmin/productos/{p.pk}/", data=base)
        self.assertEqual(r.status_code, 302)
        self.assertEqual(p.imagenes.count(), 0)

    def test_tipos_crud_y_proteccion(self):
        from catalog.models import TipoProducto

        self.client.force_login(self.staff)
        self.assertEqual(self.client.get("/webadmin/tipos/").status_code, 200)
        r = self.client.post("/webadmin/tipos/nuevo/", data={
            "nombre": "Teja Colonial", "slug": "teja-colonial",
            "descripcion": "", "orden": "9",
        })
        self.assertEqual(r.status_code, 302)
        tipo = TipoProducto.objects.get(slug="teja-colonial")
        html = self.client.get("/catalogo/").content.decode()
        self.assertIn("Teja Colonial", html)
        # Con producto no se elimina
        from catalog.models import Category, Product
        cat = Category.objects.create(nombre="T", slug="t")
        Product.objects.create(
            categoria=cat, tipo=tipo, sku="P-T", nombre="PT", slug="pt",
            precio_mayor="5", stock_cantidad="1")
        r = self.client.post(f"/webadmin/tipos/{tipo.pk}/eliminar/")
        self.assertEqual(r.status_code, 200)
        self.assertIn("No se puede eliminar", r.content.decode())
        self.assertTrue(TipoProducto.objects.filter(slug="teja-colonial").exists())

    def test_producto_con_tipo_y_badge_desde_panel(self):
        from catalog.models import Category, Product, TipoProducto

        self.client.force_login(self.staff)
        cat = Category.objects.create(nombre="Tejas", slug="tejas-p")
        tipo = TipoProducto.objects.create(nombre="Roma", slug="roma-p")
        r = self.client.post("/webadmin/productos/nuevo/", data={
            "categoria": cat.pk, "tipo": tipo.pk, "sku": "MPG-TP",
            "nombre": "Teja panel", "slug": "teja-panel", "descripcion": "",
            "badge": "MÁS VENDIDO", "precio_mayor": "20.00", "unidad": "m2",
            "stock_cantidad": "3", "origen": "", "imagen_url": "",
            "en_stock": "on", "activo": "on",
        })
        self.assertEqual(r.status_code, 302)
        p = Product.objects.get(sku="MPG-TP")
        self.assertEqual(p.tipo, tipo)
        self.assertEqual(p.badge, "MÁS VENDIDO")

    def test_categoria_con_productos_no_se_elimina(self):
        from catalog.models import Category, Product

        self.client.force_login(self.staff)
        cat = Category.objects.create(nombre="Bebidas", slug="bebidas")
        Product.objects.create(
            categoria=cat, sku="P-1", nombre="Piso test", slug="piso-test",
            precio_mayor="10.00", stock_cantidad="5")
        r = self.client.post(f"/webadmin/categorias/{cat.pk}/eliminar/")
        self.assertEqual(r.status_code, 200)
        self.assertIn("No se puede eliminar", r.content.decode())
        self.assertTrue(Category.objects.filter(slug="bebidas").exists())

    def test_staff_crea_producto_desde_panel(self):
        from catalog.models import Category, Product

        cat = Category.objects.create(nombre="Porcelanatos", slug="porcelanato")
        self.client.force_login(self.staff)
        resp = self.client.post("/webadmin/productos/nuevo/", data={
            "categoria": cat.pk, "sku": "MPG-PANEL", "nombre": "Panel test",
            "slug": "panel-test", "descripcion": "", "precio_mayor": "10.00",
            "unidad": "m2", "stock_cantidad": "5", "origen": "", "imagen_url": "",
            "en_stock": "on", "activo": "on",
        })
        self.assertEqual(resp.status_code, 302)
        self.assertTrue(Product.objects.filter(sku="MPG-PANEL").exists())

    def test_lista_productos_filtra_por_categoria(self):
        from catalog.models import Category, Product

        tejas = Category.objects.create(nombre="TejasX", slug="tejas-x")
        pisos = Category.objects.create(nombre="PisosX", slug="pisos-x")
        Product.objects.create(categoria=tejas, sku="T-1", nombre="Teja Uno",
                               slug="teja-uno", precio_mayor="10", stock_cantidad="1")
        Product.objects.create(categoria=pisos, sku="P-1", nombre="Piso Uno",
                               slug="piso-uno", precio_mayor="10", stock_cantidad="1")
        self.client.force_login(self.staff)

        html = self.client.get("/webadmin/productos/?categoria=tejas-x").content.decode()
        self.assertIn("Teja Uno", html)
        self.assertNotIn("Piso Uno", html)
        self.assertIn('aria-current="page"', html)  # categoría activa marcada
        self.assertIn("PisosX", html)  # las otras categorías siguen accesibles

        html = self.client.get("/webadmin/productos/?categoria=tejas-x&q=Piso").content.decode()
        self.assertNotIn("Teja Uno", html)
        self.assertNotIn("Piso Uno", html)

        html = self.client.get("/webadmin/productos/").content.decode()
        self.assertIn("Teja Uno", html)
        self.assertIn("Piso Uno", html)

    def test_alta_producto_dentro_de_categoria(self):
        from catalog.models import Category, Product

        cat = Category.objects.create(nombre="GriferíaX", slug="griferia-x")
        self.client.force_login(self.staff)

        html = self.client.get("/webadmin/productos/nuevo/?categoria=griferia-x").content.decode()
        self.assertIn('name="categoria" value="%d"' % cat.pk, html)
        self.assertIn('name="volver_categoria" value="griferia-x"', html)
        self.assertNotIn('<select name="categoria"', html)  # categoría fijada, sin select

        r = self.client.post("/webadmin/productos/nuevo/?categoria=griferia-x", data={
            "categoria": cat.pk, "volver_categoria": "griferia-x",
            "sku": "GR-1", "nombre": "Grifería Uno", "slug": "griferia-uno",
            "descripcion": "", "precio_mayor": "12.50", "unidad": "unid",
            "stock_cantidad": "4", "origen": "", "imagen_url": "https://example.com/y.jpg",
            "en_stock": "on", "activo": "on",
        })
        self.assertEqual(r.status_code, 302)
        self.assertIn("?categoria=griferia-x", r["Location"])
        p = Product.objects.get(sku="GR-1")
        self.assertEqual(p.categoria, cat)

    def test_editar_producto_dentro_de_categoria(self):
        from catalog.models import Category, Product

        cat = Category.objects.create(nombre="GriferíaY", slug="griferia-y")
        p = Product.objects.create(categoria=cat, sku="GR-2", nombre="Grifería Dos",
                                   slug="griferia-dos", precio_mayor="9.99", unidad="unid")
        self.client.force_login(self.staff)

        r = self.client.get("/webadmin/productos/%d/?categoria=griferia-y" % p.pk)
        self.assertEqual(r.status_code, 200)
        self.assertIn('name="volver_categoria" value="griferia-y"', r.content.decode())

        r = self.client.post("/webadmin/productos/%d/?categoria=griferia-y" % p.pk, data={
            "categoria": cat.pk, "volver_categoria": "griferia-y",
            "sku": "GR-2", "nombre": "Grifería Dos Editada", "slug": "griferia-dos",
            "descripcion": "", "precio_mayor": "10.50", "unidad": "unid",
            "stock_cantidad": "0", "origen": "", "imagen_url": "",
            "activo": "on",
        })
        self.assertEqual(r.status_code, 302)
        self.assertIn("?categoria=griferia-y", r["Location"])
        p.refresh_from_db()
        self.assertEqual(p.nombre, "Grifería Dos Editada")
        self.assertEqual(str(p.precio_mayor), "10.50")

    def test_alta_rapida_categoria_y_tipo(self):
        from catalog.models import Category, TipoProducto

        self.client.force_login(self.staff)

        # Sólo con el nombre: slug y orden se completan solos
        r = self.client.post("/webadmin/categorias/rapida/", data={"nombre": "Nueva Cat X"})
        self.assertEqual(r.status_code, 200)
        body = r.json()
        self.assertTrue(body["ok"])
        cat = Category.objects.get(slug="nueva-cat-x")
        self.assertEqual(body["id"], cat.pk)
        self.assertEqual(body["nombre"], "Nueva Cat X")

        r = self.client.post("/webadmin/tipos/rapido/", data={"nombre": "Nuevo Tipo X"})
        self.assertEqual(r.status_code, 200)
        self.assertTrue(r.json()["ok"])
        self.assertTrue(TipoProducto.objects.filter(slug="nuevo-tipo-x").exists())

        # Slug duplicado -> 400 con errores legibles
        r = self.client.post("/webadmin/categorias/rapida/", data={"nombre": "Nueva Cat X"})
        self.assertEqual(r.status_code, 400)
        body = r.json()
        self.assertFalse(body["ok"])
        self.assertIn("slug", body["errors"])
        self.assertEqual(body["errors"]["slug"],
                         ["Ya existe un registro con ese nombre."])

    def test_alta_rapida_requiere_staff(self):
        from catalog.models import Category

        self.client.force_login(self.plain)
        r = self.client.post("/webadmin/categorias/rapida/", data={"nombre": "X"})
        self.assertNotEqual(r.status_code, 200)
        self.assertFalse(Category.objects.filter(nombre="X").exists())

    def test_formulario_producto_tiene_altas_rapidas(self):
        self.client.force_login(self.staff)

        html = self.client.get("/webadmin/productos/nuevo/").content.decode()
        self.assertIn('data-alta-rapida="categoria"', html)
        self.assertIn('data-alta-rapida="tipo"', html)
        self.assertIn("/webadmin/categorias/rapida/", html)
        self.assertIn("/webadmin/tipos/rapido/", html)
        self.assertIn("<dialog", html)
        self.assertIn('name="csrfmiddlewaretoken"', html)

        # Con la categoría fijada (alta desde el listado) no se ofrece crear categoría.
        from catalog.models import Category

        Category.objects.create(nombre="Cat Modal", slug="cat-modal")
        html = self.client.get("/webadmin/productos/nuevo/?categoria=cat-modal").content.decode()
        self.assertNotIn('data-alta-rapida="categoria"', html)
        self.assertIn('data-alta-rapida="tipo"', html)

    def _excel_productos(self, filas):
        from io import BytesIO

        from openpyxl import load_workbook

        from catalog.excel_productos import build_productos_xlsx

        workbook = load_workbook(BytesIO(build_productos_xlsx()))
        sheet = workbook.active
        for fila in filas:
            sheet.append(fila)
        output = BytesIO()
        workbook.save(output)
        return output.getvalue()

    def _subir(self, contenido, nombre="productos.xlsx", categoria=None):
        from django.core.files.uploadedfile import SimpleUploadedFile

        data = {"archivo": SimpleUploadedFile(
            nombre, contenido,
            content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")}
        if categoria:
            data["categoria"] = categoria
        return self.client.post("/webadmin/productos/importar/", data=data)

    def test_plantilla_productos_descarga_y_requiere_staff(self):
        from io import BytesIO

        from openpyxl import load_workbook

        from catalog.excel_productos import HEADERS

        r = self.client.get("/webadmin/productos/plantilla/")
        self.assertNotEqual(r.status_code, 200)

        self.client.force_login(self.staff)
        r = self.client.get("/webadmin/productos/plantilla/")
        self.assertEqual(r.status_code, 200)
        self.assertIn("spreadsheet", r["Content-Type"])
        self.assertIn(".xlsx", r["Content-Disposition"])
        workbook = load_workbook(BytesIO(r.content), read_only=True)
        self.assertEqual([c.value for c in workbook.active[1]], HEADERS)

    def test_importar_excel_crea_actualiza_y_usa_categoria_por_defecto(self):
        from catalog.models import Category, Product

        cat = Category.objects.create(nombre="Tejas Imp", slug="tejas-imp")
        Product.objects.create(categoria=cat, sku="IMP-2", nombre="Viejo",
                               slug="imp-2", precio_mayor="1.00", unidad="m2")
        self.client.force_login(self.staff)

        contenido = self._excel_productos([
            ["IMP-1", "Nuevo Uno", "Tejas Imp", None, 20, "unid", 5, "", "", "", ""],
            ["IMP-2", "Renombrado", "Tejas Imp", None, 3.5, None, None, "", "", "", ""],
            ["IMP-3", "Sin categoría", None, None, 1, None, None, None, None, None, None],
        ])
        r = self._subir(contenido, categoria="tejas-imp")
        self.assertEqual(r.status_code, 200)
        html = r.content.decode()
        self.assertIn("2 creados", html)
        self.assertIn("1 actualizado", html)

        self.assertEqual(Product.objects.get(sku="IMP-1").slug, "nuevo-uno")
        self.assertEqual(str(Product.objects.get(sku="IMP-1").precio_mayor), "20.00")
        self.assertEqual(Product.objects.get(sku="IMP-1").unidad, "unid")
        actualizado = Product.objects.get(sku="IMP-2")
        self.assertEqual(actualizado.nombre, "Renombrado")
        self.assertEqual(str(actualizado.precio_mayor), "3.50")
        self.assertEqual(Product.objects.get(sku="IMP-3").categoria, cat)

    def test_importar_excel_sin_sku_asigna_el_correlativo(self):
        """Si el Excel deja el SKU vacío, el importador genera COD-0001…"""
        from catalog.models import Category, Product

        cat = Category.objects.create(nombre="Tejas SKU", slug="tejas-sku")
        Product.objects.create(categoria=cat, sku="COD-0001", nombre="Existente",
                               slug="existente", precio_mayor="1.00", unidad="m2")
        self.client.force_login(self.staff)

        contenido = self._excel_productos([
            ["", "Sin SKU A", "Tejas SKU", None, 20, "unid", 5, "", "", "", ""],
            ["", "Sin SKU B", "Tejas SKU", None, 30, "unid", 5, "", "", "", ""],
        ])
        r = self._subir(contenido, categoria="tejas-sku")
        self.assertEqual(r.status_code, 200)
        self.assertIn("2 creados", r.content.decode())
        self.assertTrue(Product.objects.filter(sku="COD-0002").exists())
        self.assertTrue(Product.objects.filter(sku="COD-0003").exists())
        self.assertEqual(Product.objects.get(sku="COD-0002").nombre, "Sin SKU A")

    def test_importar_excel_con_colores_y_galeria(self):
        """El Excel puede traer los colores del modelo y las fotos de la galería."""
        from catalog.models import Category, Product

        cat = Category.objects.create(nombre="Tejas Col", slug="tejas-col")
        self.client.force_login(self.staff)

        contenido = self._excel_productos([
            ["EXC-1", "Teja con colores", "Tejas Col", None, 20, "m2", 5, "", "",
             "https://ejemplo.com/portada.jpg", "Descripción",
             "Terracota, Negro (exterior), Blanco (interior)",
             "https://ejemplo.com/f1.jpg;https://ejemplo.com/f2.jpg"],
        ])
        r = self._subir(contenido, categoria="tejas-col")
        self.assertEqual(r.status_code, 200)
        p = Product.objects.get(sku="EXC-1")
        colores = list(p.colores_producto.all())
        self.assertEqual([c.color.nombre for c in colores],
                         ["Terracota", "Negro", "Blanco"])
        self.assertEqual([c.rol for c in colores], ["", "exterior", "interior"])
        urls = [f.imagen_url for f in p.imagenes.all()]
        self.assertEqual(urls, ["https://ejemplo.com/f1.jpg",
                                "https://ejemplo.com/f2.jpg"])

    def test_importar_excel_no_pisa_productos_con_sku_generado(self):
        """Un SKU autogenerado no puede pisar el SKU que trae otra fila."""
        from catalog.models import Category, Product

        cat = Category.objects.create(nombre="Tejas SKU2", slug="tejas-sku2")
        self.client.force_login(self.staff)
        contenido = self._excel_productos([
            ["", "Sin SKU A", "Tejas SKU2", None, 20, "m2", 5, "", "", "", ""],
            ["COD-0001", "Con SKU B", "Tejas SKU2", None, 30, "m2", 5, "", "", "", ""],
        ])
        r = self._subir(contenido, categoria="tejas-sku2")
        self.assertEqual(r.status_code, 200)
        # Ambos productos existen: el generado toma un número libre.
        self.assertTrue(Product.objects.filter(nombre="Sin SKU A").exists())
        self.assertTrue(Product.objects.filter(sku="COD-0001",
                                               nombre="Con SKU B").exists())

    def test_importar_excel_con_color_repetido_no_rompe(self):
        from catalog.models import Category, Product

        cat = Category.objects.create(nombre="Tejas Dup", slug="tejas-dup")
        self.client.force_login(self.staff)
        contenido = self._excel_productos([
            ["DUP-COL", "Teja colores repetidos", "Tejas Dup", None, 20, "m2", 5,
             "", "", "", "", "Negro, Negro, negro", ""],
        ])
        r = self._subir(contenido, categoria="tejas-dup")
        self.assertEqual(r.status_code, 200)
        p = Product.objects.get(sku="DUP-COL")
        self.assertEqual(p.colores_producto.count(), 1)

    def test_galeria_rechaza_archivos_que_no_son_imagenes(self):
        from io import BytesIO

        from django.core.files.uploadedfile import SimpleUploadedFile

        from catalog.models import Category, Product

        self.client.force_login(self.staff)
        cat = Category.objects.create(nombre="Tejas Mal", slug="tejas-mal")
        archivo = SimpleUploadedFile("x.html", b"<script>alert(1)</script>",
                                     content_type="text/html")
        base = {"categoria": cat.pk, "sku": "MPG-MAL", "nombre": "No imagen",
                "slug": "no-imagen", "descripcion": "", "precio_mayor": "10.00",
                "unidad": "m2", "stock_cantidad": "0", "imagen_url": "",
                "en_stock": "on", "activo": "on",
                "colores_json": "[]", "caracteristicas_json": "[]",
                "galeria_archivos": archivo}
        r = self.client.post("/webadmin/productos/nuevo/", data=base)
        self.assertEqual(r.status_code, 302)
        p = Product.objects.get(sku="MPG-MAL")
        self.assertEqual(p.imagenes.count(), 0)

    def test_fusionar_solo_acepta_el_mismo_grupo(self):
        """No se puede fusionar un producto de otra categoría."""
        from catalog.models import Category, Product

        cat = Category.objects.create(nombre="Cubiertas F", slug="cubiertas-f")
        otra = Category.objects.create(nombre="Pisos F", slug="pisos-f")
        principal = Product.objects.create(categoria=cat, sku="FUS-1",
                                           nombre="Teja F Negro/Blanco 2.5 mm",
                                           slug="fus-1", precio_mayor="10")
        Product.objects.create(categoria=cat, sku="FUS-2",
                               nombre="Teja F Verde/Blanco 2.5 mm", slug="fus-2",
                               precio_mayor="10")
        intruso = Product.objects.create(categoria=otra, sku="FUS-3", nombre="Piso F",
                                         slug="fus-3", precio_mayor="10")
        self.client.force_login(self.staff)
        r = self.client.post("/webadmin/productos/fusionar/", {
            "principal": principal.pk, "variantes": [intruso.pk]})
        self.assertEqual(r.status_code, 302)
        intruso.refresh_from_db()
        self.assertTrue(intruso.activo)  # no se toca lo que no es del grupo

    def test_fusionar_junta_el_grupo_del_mismo_modelo(self):
        from catalog.models import Category, Color, Product, ProductColor

        cat = Category.objects.create(nombre="Cubiertas G", slug="cubiertas-g")
        verde = Color.objects.get(slug="verde")
        principal = Product.objects.create(categoria=cat, sku="FUS-10",
                                           nombre="Teja G Negro/Blanco 2.5 mm",
                                           slug="fus-10", precio_mayor="10")
        variante = Product.objects.create(categoria=cat, sku="FUS-11",
                                          nombre="Teja G Verde/Blanco 2.5 mm",
                                          slug="fus-11", precio_mayor="10")
        ProductColor.objects.create(producto=variante, color=verde, orden=0)
        self.client.force_login(self.staff)
        r = self.client.post("/webadmin/productos/fusionar/", {
            "principal": principal.pk, "variantes": [variante.pk]})
        self.assertEqual(r.status_code, 302)
        variante.refresh_from_db()
        self.assertFalse(variante.activo)
        # El color de la variante pasa al principal.
        self.assertEqual([pc.color.slug for pc in principal.colores_producto.all()],
                         ["verde"])

    def test_importar_excel_errores_y_permiso(self):
        from catalog.models import Product

        contenido = self._excel_productos([["IMP-9", "X", "No Existe", None, 1, None, None, None, None, None, None]])

        self.client.force_login(self.plain)
        r = self._subir(contenido)
        self.assertNotEqual(r.status_code, 200)
        self.assertEqual(Product.objects.count(), 0)

        self.client.force_login(self.staff)
        r = self._subir(b"no soy un excel", nombre="malo.xlsx")
        self.assertEqual(r.status_code, 200)
        self.assertIn("No se pudo leer el archivo Excel", r.content.decode())

        r = self._subir(contenido)
        self.assertEqual(r.status_code, 200)
        self.assertIn("Categoría no encontrada", r.content.decode())

        r = self._subir(b"0" * (5 * 1024 * 1024 + 1), nombre="grande.xlsx")
        self.assertEqual(r.status_code, 200)
        self.assertIn("5 MB", r.content.decode())

        self.assertEqual(Product.objects.count(), 0)

    def _producto(self, sku, nombre, categoria, **kwargs):
        from catalog.models import Product

        datos = {"precio_mayor": kwargs.pop("precio_mayor", "10.00"),
                 "stock_cantidad": kwargs.pop("stock_cantidad", "1.00")}
        datos.update(kwargs)
        return Product.objects.create(
            categoria=categoria, sku=sku, nombre=nombre, slug=sku.lower(), **datos)

    def _categorias_acciones(self):
        from catalog.models import Category

        origen = Category.objects.create(nombre="AccOrigen", slug="acc-origen")
        destino = Category.objects.create(nombre="AccDestino", slug="acc-destino")
        return origen, destino

    def test_precios_inline_actualiza_precio_y_stock(self):
        origen, _ = self._categorias_acciones()
        p1 = self._producto("ACC-1", "Uno", origen)
        p2 = self._producto("ACC-2", "Dos", origen, stock_cantidad="3.00")
        self.client.force_login(self.staff)

        r = self.client.post("/webadmin/productos/precios/", data={
            "ids_precio": [p1.pk, p2.pk],
            "precio_%d" % p1.pk: "15.5", "stock_%d" % p1.pk: "7",
            "precio_%d" % p2.pk: "0", "stock_%d" % p2.pk: "2,5",
            "volver_categoria": "acc-origen"})
        self.assertEqual(r.status_code, 302)
        self.assertIn("categoria=acc-origen", r["Location"])
        self.assertIn("msg=precios", r["Location"])
        p1.refresh_from_db()
        p2.refresh_from_db()
        self.assertEqual(str(p1.precio_mayor), "15.50")
        self.assertEqual(str(p1.stock_cantidad), "7.00")
        self.assertEqual(str(p2.precio_mayor), "0.00")
        self.assertEqual(str(p2.stock_cantidad), "2.50")

    def test_precios_inline_rechaza_datos_invalidos(self):
        origen, _ = self._categorias_acciones()
        p1 = self._producto("ACC-1", "Uno", origen)
        self.client.force_login(self.staff)

        for valor in ("abc", "-1", ""):
            with self.subTest(valor=valor):
                r = self.client.post("/webadmin/productos/precios/", data={
                    "ids_precio": [p1.pk],
                    "precio_%d" % p1.pk: valor, "stock_%d" % p1.pk: "1"})
                self.assertEqual(r.status_code, 302)
                self.assertIn("msg=datos_invalidos", r["Location"])
                p1.refresh_from_db()
                self.assertEqual(str(p1.precio_mayor), "10.00")

    def test_precios_inline_requiere_staff(self):
        origen, _ = self._categorias_acciones()
        p1 = self._producto("ACC-1", "Uno", origen)
        self.client.force_login(self.plain)
        r = self.client.post("/webadmin/productos/precios/", data={
            "ids_precio": [p1.pk],
            "precio_%d" % p1.pk: "99", "stock_%d" % p1.pk: "9"})
        self.assertNotEqual(r.status_code, 200)
        p1.refresh_from_db()
        self.assertEqual(str(p1.precio_mayor), "10.00")

    def test_badges_de_faltantes_y_conteo_por_categoria(self):
        origen, _ = self._categorias_acciones()
        self._producto("CAL-OK", "Completo", origen,
                       imagen_url="https://example.com/a.jpg",
                       caracteristicas=["Resistente UV"], precio_mayor="12.00")
        self._producto("CAL-VAC", "Vacío", origen, precio_mayor="0.00")
        self.client.force_login(self.staff)

        html = self.client.get("/webadmin/productos/").content.decode()
        # Sólo el producto vacío los lleva (una vez cada uno)
        self.assertEqual(html.count("sin fotos"), 1)
        self.assertEqual(html.count("sin características"), 1)
        self.assertEqual(html.count("sin precio"), 1)
        # El pill de la categoría avisa de 1 ficha incompleta
        self.assertEqual(html.count('data-incompletos="1"'), 1)
        self.assertEqual(html.count("data-incompletos="), 1)

    def test_busqueda_global_con_todas(self):
        origen, destino = self._categorias_acciones()
        self._producto("G-1", "Encontrable Uno", origen)
        self._producto("G-2", "Encontrable Dos", destino)
        self.client.force_login(self.staff)

        html = self.client.get(
            "/webadmin/productos/?categoria=acc-origen&q=Encontrable").content.decode()
        self.assertIn("Encontrable Uno", html)
        self.assertNotIn("Encontrable Dos", html)

        html = self.client.get(
            "/webadmin/productos/?categoria=acc-origen&q=Encontrable&todas=1").content.decode()
        self.assertIn("Encontrable Uno", html)
        self.assertIn("Encontrable Dos", html)
        self.assertIn('name="todas"', html)
        self.assertIn("checked", html)

    def test_aviso_de_accion_en_la_lista(self):
        origen, _ = self._categorias_acciones()
        self._producto("AV-1", "Uno", origen)
        self.client.force_login(self.staff)

        html = self.client.get("/webadmin/productos/?msg=eliminados&n=2").content.decode()
        self.assertIn("2 productos eliminados", html)

        html = self.client.get("/webadmin/productos/?msg=datos_invalidos").content.decode()
        self.assertIn("Revisa precio y stock", html)

    def test_lista_es_tabla_con_precios_y_sin_acciones_masivas(self):
        origen, _ = self._categorias_acciones()
        p1 = self._producto("MK-1", "Uno", origen)
        self.client.force_login(self.staff)

        html = self.client.get("/webadmin/productos/").content.decode()
        self.assertIn("<table", html)
        self.assertIn("<caption", html)
        self.assertNotIn('id="acciones-masivas"', html)
        self.assertIn('id="precios"', html)
        self.assertIn('form="precios"', html)
        self.assertIn('name="ids_precio"', html)
        self.assertIn('name="precio_%d"' % p1.pk, html)
        self.assertIn('name="stock_%d"' % p1.pk, html)
        self.assertIn("data-editar-producto", html)
        self.assertIn('id="dialogo-editar"', html)
        self.assertIn('id="dialogo-excel"', html)
        self.assertIn('id="alta-rapida"', html)

    def test_modal_de_edicion_devuelve_solo_el_formulario(self):
        origen, _ = self._categorias_acciones()
        p1 = self._producto("MOD-1", "Uno", origen)
        self.client.force_login(self.staff)

        pagina = self.client.get("/webadmin/productos/%d/" % p1.pk).content.decode()
        self.assertIn("<!DOCTYPE html>", pagina)

        r = self.client.get(
            "/webadmin/productos/%d/?modal=1&categoria=acc-origen" % p1.pk)
        self.assertEqual(r.status_code, 200)
        html = r.content.decode()
        self.assertNotIn("<!DOCTYPE html>", html)
        self.assertIn('id="id_sku"', html)
        self.assertIn('action="/webadmin/productos/%d/?categoria=acc-origen"' % p1.pk, html)

    def test_precios_y_stock_sin_localizar_en_inputs_number(self):
        origen, _ = self._categorias_acciones()
        p1 = self._producto("LOC-1", "Uno", origen,
                             precio_mayor="1234.50", stock_cantidad="99.25")
        self.client.force_login(self.staff)

        html = self.client.get("/webadmin/productos/").content.decode()
        self.assertIn('name="precio_%d" value="1234.50"' % p1.pk, html)
        self.assertIn('name="stock_%d" value="99.25"' % p1.pk, html)
        self.assertNotIn("%d\" value=\"1234,50\"" % p1.pk, html)

    def test_pagina_duplicados_agrupa_y_requiere_staff(self):
        origen, destino = self._categorias_acciones()
        p1 = self._producto("DUP-1", "Teja Roma PVC", origen)
        p2 = self._producto("DUP-2", "TEJA  roma pvc", origen)
        self._producto("DUP-3", "Teja Roma PVC", destino)

        self.client.force_login(self.plain)
        self.assertNotEqual(self.client.get("/webadmin/productos/duplicados/").status_code, 200)

        self.client.force_login(self.staff)
        html = self.client.get("/webadmin/productos/duplicados/").content.decode()
        self.assertIn("DUP-1", html)
        self.assertIn("DUP-2", html)
        self.assertNotIn("DUP-3", html)  # en otra categoría: no es duplicado
        self.assertIn("1 grupo repetido", html)
        self.assertIn("2 productos", html)
        self.assertIn("/webadmin/productos/%d/" % p1.pk, html)
        self.assertIn("/webadmin/productos/%d/eliminar/" % p2.pk, html)

    def test_producto_eliminar_confirma_y_borra(self):
        from catalog.models import Product

        origen, _ = self._categorias_acciones()
        p1 = self._producto("DEL-1", "Uno", origen)
        self.client.force_login(self.staff)

        r = self.client.get("/webadmin/productos/%d/eliminar/" % p1.pk)
        self.assertEqual(r.status_code, 200)
        self.assertIn("DEL-1", r.content.decode())
        self.assertTrue(Product.objects.filter(pk=p1.pk).exists())

        r = self.client.post("/webadmin/productos/%d/eliminar/" % p1.pk,
                             data={"volver_categoria": "acc-origen"})
        self.assertEqual(r.status_code, 302)
        self.assertIn("categoria=acc-origen", r["Location"])
        self.assertIn("msg=eliminados", r["Location"])
        self.assertFalse(Product.objects.filter(pk=p1.pk).exists())

    def test_producto_eliminar_requiere_staff(self):
        from catalog.models import Product

        origen, _ = self._categorias_acciones()
        p1 = self._producto("DEL-2", "Dos", origen)
        self.client.force_login(self.plain)
        self.assertNotEqual(
            self.client.get("/webadmin/productos/%d/eliminar/" % p1.pk).status_code, 200)
        self.client.post("/webadmin/productos/%d/eliminar/" % p1.pk)
        self.assertTrue(Product.objects.filter(pk=p1.pk).exists())


class ColaboradoresTest(TestCase):
    def setUp(self):
        from django.contrib.auth import get_user_model

        from panel.models import AccesoPanel

        self.AccesoPanel = AccesoPanel
        self.User = get_user_model()
        self.dueno = self.User.objects.create_user("dueno", password="x", is_staff=True)
        self.empleado = self.User.objects.create_user("calc", password="clave-larga-1")
        AccesoPanel.objects.create(usuario=self.empleado, calculadora=True)

    def _perfil(self, **cambios):
        perfil = self.AccesoPanel.objects.get(usuario=self.empleado)
        for clave, valor in cambios.items():
            setattr(perfil, clave, valor)
        perfil.save()
        return perfil

    def test_empleado_aterriza_en_su_seccion(self):
        self.client.force_login(self.empleado)
        r = self.client.get("/webadmin/")
        self.assertRedirects(r, "/webadmin/calculadora/")
        self.assertEqual(self.client.get("/webadmin/calculadora/").status_code, 200)

    def test_empleado_recibe_403_fuera_de_sus_secciones(self):
        self.client.force_login(self.empleado)
        for url in ("/webadmin/productos/", "/webadmin/pedidos/",
                    "/webadmin/codigos/", "/webadmin/productos/duplicados/"):
            self.assertEqual(self.client.get(url).status_code, 403, url)
        for url in ("/webadmin/ajustes/", "/webadmin/colaboradores/"):
            self.assertEqual(self.client.get(url).status_code, 302, url)

    def test_nav_del_empleado_solo_muestra_lo_permitido(self):
        self.client.force_login(self.empleado)
        html = self.client.get("/webadmin/calculadora/").content.decode()
        self.assertIn('href="/webadmin/calculadora/"', html)
        self.assertNotIn('href="/webadmin/productos/"', html)
        self.assertNotIn('href="/webadmin/pedidos/"', html)
        self.assertNotIn('href="/webadmin/colaboradores/"', html)
        self.assertNotIn('href="/webadmin/ajustes/"', html)
        self.assertNotIn("Tablero</a>", html)
        self.assertIn("Colaborador", html)

    def test_dueno_ve_todo_y_colaboradores(self):
        self.client.force_login(self.dueno)
        html = self.client.get("/webadmin/").content.decode()
        for url in ("/webadmin/productos/", "/webadmin/pedidos/", "/webadmin/codigos/",
                    "/webadmin/calculadora/", "/webadmin/ajustes/",
                    "/webadmin/colaboradores/"):
            self.assertIn('href="%s"' % url, html)
        self.assertEqual(self.client.get("/webadmin/colaboradores/").status_code, 200)

    def test_empleado_con_productos_entra_al_catalogo(self):
        self._perfil(calculadora=False, productos=True)
        self.client.force_login(self.empleado)
        r = self.client.get("/webadmin/")
        self.assertRedirects(r, "/webadmin/productos/")
        self.assertEqual(self.client.get("/webadmin/productos/").status_code, 200)
        self.assertEqual(self.client.get("/webadmin/calculadora/").status_code, 403)

    def test_empleado_desactivado_no_puede_entrar(self):
        self._perfil(activo=False)
        r = self.client.post("/webadmin/login/",
                             {"username": "calc", "password": "clave-larga-1"})
        self.assertEqual(r.status_code, 200)
        self.assertContains(r, "no tiene acceso al panel")

    def test_usuario_sin_perfil_no_puede_entrar(self):
        self.User.objects.create_user("cliente", password="clave-larga-1")
        r = self.client.post("/webadmin/login/",
                             {"username": "cliente", "password": "clave-larga-1"})
        self.assertEqual(r.status_code, 200)
        self.assertContains(r, "no tiene acceso al panel")

    def test_dueno_crea_colaborador_con_secciones(self):
        self.client.force_login(self.dueno)
        r = self.client.post("/webadmin/colaboradores/nuevo/", {
            "username": "bodega", "first_name": "Bodega",
            "password": "ClaveSegura-2026", "productos": "on", "activo": "on"})
        self.assertRedirects(r, "/webadmin/colaboradores/")
        usuario = self.User.objects.get(username="bodega")
        self.assertFalse(usuario.is_staff)
        self.assertTrue(usuario.check_password("ClaveSegura-2026"))
        perfil = usuario.acceso_panel
        self.assertTrue(perfil.productos)
        self.assertFalse(perfil.calculadora)
        self.assertTrue(perfil.activo)

    def test_exige_al_menos_una_seccion(self):
        self.client.force_login(self.dueno)
        r = self.client.post("/webadmin/colaboradores/nuevo/", {
            "username": "sinseccion", "password": "ClaveSegura-2026", "activo": "on"})
        self.assertEqual(r.status_code, 200)
        self.assertContains(r, "Marca al menos una sección")
        self.assertFalse(self.User.objects.filter(username="sinseccion").exists())

    def test_usuario_duplicado(self):
        self.client.force_login(self.dueno)
        r = self.client.post("/webadmin/colaboradores/nuevo/", {
            "username": "CALC", "password": "ClaveSegura-2026",
            "calculadora": "on", "activo": "on"})
        self.assertEqual(r.status_code, 200)
        self.assertContains(r, "Ya existe un usuario con ese nombre")

    def test_editar_sin_password_conserva_la_clave(self):
        perfil = self._perfil()
        self.client.force_login(self.dueno)
        r = self.client.post("/webadmin/colaboradores/%d/" % perfil.pk, {
            "username": "calc", "first_name": "Calc",
            "productos": "on", "activo": "on"})
        self.assertRedirects(r, "/webadmin/colaboradores/")
        self.empleado.refresh_from_db()
        self.assertTrue(self.empleado.check_password("clave-larga-1"))
        perfil.refresh_from_db()
        self.assertTrue(perfil.productos)
        self.assertFalse(perfil.calculadora)

    def test_desactivar_y_eliminar(self):
        perfil = self._perfil()
        self.client.force_login(self.dueno)
        self.client.post("/webadmin/colaboradores/%d/estado/" % perfil.pk)
        perfil.refresh_from_db()
        self.empleado.refresh_from_db()
        self.assertFalse(perfil.activo)
        self.assertFalse(self.empleado.is_active)
        self.client.post("/webadmin/colaboradores/%d/eliminar/" % perfil.pk)
        self.assertFalse(self.AccesoPanel.objects.filter(pk=perfil.pk).exists())
        self.assertFalse(self.User.objects.filter(username="calc").exists())

    def test_el_dueno_no_se_edita_desde_colaboradores(self):
        self.client.force_login(self.dueno)
        r = self.client.get("/webadmin/colaboradores/%d/" % self.dueno.pk)
        self.assertEqual(r.status_code, 404)
